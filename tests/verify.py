# -*- coding: utf-8 -*-
"""校验拆分结果（需要 openpyxl，仅用于开发验证）"""
import csv
import datetime
import os
import sys

from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "测试数据", "员工工资表_拆分结果")
SRC = sys.argv[2] if len(sys.argv) > 2 else os.path.join(BASE, "测试数据", "员工工资表.xlsx")

fail = []


def check(cond, msg):
    if cond:
        print("  [OK] %s" % msg)
    else:
        print("  [FAIL] %s" % msg)
        fail.append(msg)


def main():
    print("校验输出目录：%s" % OUT)
    files = [f for f in os.listdir(OUT) if f.endswith(".xlsx")]
    print("生成文件数：%d" % len(files))
    for f in sorted(files):
        print("   - %s" % f)

    # 源数据
    wb = load_workbook(SRC, data_only=True)
    ws = wb["员工工资"]
    src_header = [c.value for c in ws[1]]
    src_rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    print("\n源表：%d 行数据" % len(src_rows))

    # 非法字符检查
    bad = [f for f in files if any(ch in f for ch in '\\/:*?"<>|')]
    check(not bad, "输出文件名不含非法字符（%s）" % (bad or "通过"))

    # 汇总校验
    total = 0
    dept_col = src_header.index("部门")
    for f in sorted(files):
        path = os.path.join(OUT, f)
        wb2 = load_workbook(path, data_only=True)
        ws2 = wb2.active
        header = [c.value for c in ws2[1]]
        check(header == src_header, "%s 表头一致" % f)
        vals = [list(r) for r in ws2.iter_rows(min_row=2, values_only=True)]
        total += len(vals)
        # 分类列取值唯一（除空值组）
        uniq = set(str(r[dept_col]).strip() if r[dept_col] is not None else "" for r in vals)
        check(len(uniq) <= 1, "%s 部门列取值唯一：%s" % (f, list(uniq)[:2]))
        # 日期类型
        dtypes = set(type(r[5]).__name__ for r in vals if r[5] is not None)
        check(dtypes <= {"datetime"}, "%s 入职日期为日期类型（%s）" % (f, dtypes or "空"))
        # 数字类型
        ntypes = set(type(r[3]).__name__ for r in vals if r[3] is not None)
        check(ntypes <= {"int"}, "%s 工资为整数类型（%s）" % (f, ntypes or "空"))
        bt = set(type(r[6]).__name__ for r in vals if r[6] is not None)
        check(bt <= {"bool"}, "%s 是否在职为布尔类型（%s）" % (f, bt or "空"))

    print("\n拆分后数据总行数：%d（源 %d）" % (total, len(src_rows)))
    check(total == len(src_rows), "行数守恒")

    # 清单校验
    sp = os.path.join(OUT, "拆分清单.csv")
    if os.path.exists(sp):
        with open(sp, encoding="utf-8-sig") as fh:
            rows = list(csv.reader(fh))
        s = sum(int(r[2]) for r in rows[1:])
        check(s == len(src_rows), "拆分清单行数合计 == 源行数（%d）" % s)
        check(len(rows) - 1 == len(files), "清单条目数 == 文件数")
    else:
        check(False, "存在拆分清单.csv")

    print("\n" + ("=" * 50))
    if fail:
        print("❌ 失败 %d 项：" % len(fail))
        for m in fail:
            print("   - %s" % m)
        return 1
    print("✅ 全部校验通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
