# -*- coding: utf-8 -*-
"""校验“沿用原表模板格式”的拆分结果（需要 openpyxl，仅开发验证用）"""
import os
import sys

from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "测试数据", "模板样表_拆分结果")
SRC = sys.argv[2] if len(sys.argv) > 2 else os.path.join(BASE, "测试数据", "模板样表.xlsx")

fail = []


def check(cond, msg, extra=""):
    print("  %s %s%s" % ("[OK]" if cond else "[FAIL]", msg, ("  → " + str(extra)) if not cond and extra else ""))
    if not cond:
        fail.append(msg)


def style(cell):
    f, fill = cell.font, cell.fill
    bd = cell.border
    return {
        "font": (f.name, f.size, f.bold, (f.color.rgb if f.color else None)),
        "fill": (fill.fgColor.rgb if fill and fill.fill_type else None),
        "border": bool(bd and (bd.left.style or bd.right.style or bd.top.style or bd.bottom.style)),
        "fmt": cell.number_format,
        "align": cell.alignment.horizontal,
    }


def main():
    print("源文件：%s\n输出目录：%s\n" % (SRC, OUT))
    src = load_workbook(SRC)
    sws = src["员工工资"]

    files = sorted(f for f in os.listdir(OUT) if f.endswith(".xlsx"))
    print("生成 %d 个文件：%s\n" % (len(files), "、".join(files)))

    target = "财务部.xlsx" if "财务部.xlsx" in files else files[0]
    print("—— 以《%s》为例逐项比对模板格式 ——" % target)
    wb = load_workbook(os.path.join(OUT, target))
    ws = wb.active

    # 1. 标题行（原第 1 行）
    check(ws["A1"].value == sws["A1"].value, "标题文字一致", ws["A1"].value)
    check(str(ws.merged_cells.ranges).find("A1:I1") >= 0, "标题行合并 A1:I1 保留",
          list(ws.merged_cells.ranges))
    s1, o1 = style(ws["A1"]), style(sws["A1"])
    check(s1["font"] == o1["font"], "标题字体一致 %s" % (s1["font"],), (s1["font"], o1["font"]))
    check(s1["fill"] == o1["fill"], "标题底色一致 %s" % (s1["fill"],), (s1["fill"], o1["fill"]))
    check(ws.row_dimensions[1].height == sws.row_dimensions[1].height,
          "标题行高一致 %s" % ws.row_dimensions[1].height, ws.row_dimensions[1].height)

    # 2. 表头行（原第 2 行）
    for col in "ABCDEFGH":
        s, o = style(ws[col + "2"]), style(sws[col + "2"])
        check(s == o, "表头 %s2 样式（字体/底色/边框/对齐）完全一致" % col, (s, o))
    check(ws["B2"].value == "姓名", "表头文字正确", ws["B2"].value)

    # 3. 数据行样式（取最后一行数据，它一定被移动过）
    last = ws.max_row
    check(last > 2, "存在数据行（最后一行=%d）" % last)
    s, o = style(ws["D%d" % last]), style(sws["D201"])
    check(s["fmt"] == o["fmt"] == "#,##0", "工资列数字格式保留 %s" % s["fmt"], s["fmt"])
    s, o = style(ws["F%d" % last]), style(sws["F201"])
    check(s["fmt"] == o["fmt"] == "yyyy-mm-dd", "日期列格式保留 %s" % s["fmt"], s["fmt"])
    s, o = style(ws["E%d" % last]), style(sws["E201"])
    check(s["fmt"] == o["fmt"] == "0.00", "绩效系数格式保留 %s" % s["fmt"], s["fmt"])
    sd = style(ws["C%d" % last])
    od = style(sws["C201"])
    check(sd["border"] and od["border"], "数据行边框保留")
    check(sd["fill"] == od["fill"], "数据行底色保留 %s" % sd["fill"], (sd["fill"], od["fill"]))
    check(ws.row_dimensions[last].height == sws.row_dimensions[201].height,
          "数据行行高保留 %s" % ws.row_dimensions[last].height)

    # 4. 视图设置
    check(ws.freeze_panes == "A3", "冻结窗格保留 %s" % ws.freeze_panes, ws.freeze_panes)
    check(ws.auto_filter.ref == "A2:I%d" % last, "自动筛选范围已重设为 A2:I%d" % last, ws.auto_filter.ref)

    # 5. 列宽
    widths_ok = all(
        ws.column_dimensions[c].width == sws.column_dimensions[c].width
        for c in "ABCDEFGHI" if c in sws.column_dimensions)
    check(widths_ok, "列宽全部保留",
          {c: (ws.column_dimensions[c].width, sws.column_dimensions[c].width) for c in "ABCDEFGHI"})

    # 6. 公式行号重映射
    bad = []
    for r in range(3, last + 1):
        v = ws["I%d" % r].value
        if v != "=D%d*E%d" % (r, r):
            bad.append((r, v))
    check(not bad, "公式行号已随新位置重映射（如 I3==D3*E3）", bad[:3])

    # 7. 工作表数量（只保留目标表）
    check(len(wb.sheetnames) == 1, "只保留目标工作表 %s" % wb.sheetnames, wb.sheetnames)

    # 8. 内容与行数
    total = 0
    for f in files:
        w2 = load_workbook(os.path.join(OUT, f))
        s2 = w2.active
        total += s2.max_row - 2
        vals = set(str(s2.cell(row=r, column=3).value or "").strip() for r in range(3, s2.max_row + 1))
        check(len(vals) <= 1, "%s：部门列取值唯一 %s" % (f, list(vals)[:2]))
    check(total == 200, "拆分后数据总行数=200（实际 %d）" % total, total)

    print("\n" + "=" * 56)
    if fail:
        print("❌ 失败 %d 项：" % len(fail))
        for m in fail:
            print("   - %s" % m)
        return 1
    print("✅ 模板格式全部保留，校验通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
